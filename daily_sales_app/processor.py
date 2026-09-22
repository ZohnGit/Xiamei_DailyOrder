from __future__ import annotations

import csv
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


ORDER_REQUIRED_FIELDS = [
    "order_status",
    "unit_price",
    "qty",
    "pcr_product_sku",
    "warehouse_product_title",
    "refrence_no_platform",
]

SKU_MAPPING_REQUIRED_FIELDS = [
    "pcr_product_sku",
    "运营",
]

COST_ORDER_ALIASES = ["平台订单号", "平台单号", "订单号"]
COST_SKU_ALIASES = ["SKU明细", "SKU 明细", "SKU"]
COST_TOTAL_ALIASES = ["总成本（欧元）", "总成本(欧元)", "总成本欧元", "总成本"]
COST_UNIT_ALIASES = ["单一数量成本", "单一数量成本（欧元）", "单一数量成本(欧元)", "单件成本", "单位成本"]

FILTER_STATUSES = {"冻结中", "付款未完成", "问题件", "已废弃"}
COST_MATCH_PLATFORMS = {"CD", "PM", "FNAC"}


@dataclass
class ProcessingResult:
    output_path: Path
    report_date: date
    grouped_rows: dict[str, list[dict[str, Any]]]
    unmatched_rows: list[dict[str, str]]


def clean_text(value: Any) -> str:
    return (
        str(value or "")
        .replace("\ufeff", "")
        .replace("\u00a0", "")
        .strip()
    )


def split_operator_names(value: Any) -> list[str]:
    raw = clean_text(value)
    if not raw:
        return []
    names = []
    for chunk in raw.replace(";", "；").split("；"):
        name = clean_text(chunk)
        if name and name not in names:
            names.append(name)
    return names


def normalize_key(value: Any) -> str:
    return clean_text(value).lower()


def normalize_token(value: Any) -> str:
    return "".join(clean_text(value).upper().split())


def parse_number(value: Any) -> float:
    if value is None:
        return 0.0
    text = clean_text(value).replace(",", "")
    if not text:
        return 0.0
    cleaned = "".join(ch for ch in text if ch.isdigit() or ch in ".-")
    if cleaned in {"", "-", ".", "-."}:
        return 0.0
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def read_table(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return _read_xlsx(path)
    if suffix == ".xls":
        raise RuntimeError(f"{path.name} 目前请先另存为 .xlsx 再导入。")
    raise RuntimeError(f"不支持的文件格式: {path.name}")


def read_raw_table(path: Path) -> list[list[Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv_raw(path)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        return _read_xlsx_raw(path)
    if suffix == ".xls":
        raise RuntimeError(f"{path.name} 目前请先另存为 .xlsx 再导入。")
    raise RuntimeError(f"不支持的文件格式: {path.name}")


def read_order_tables(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw_rows = _read_csv_raw(path)
        return _extract_order_rows_from_raw(raw_rows, path.name)
    if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        collected: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            raw_rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            collected.extend(_extract_order_rows_from_raw(raw_rows, f"{path.name}:{sheet.title}", allow_empty=True))
        if collected:
            return collected
        raise RuntimeError(f"{path.name} 没有识别到订单表头。")
    if suffix == ".xls":
        raise RuntimeError(f"{path.name} 目前请先另存为 .xlsx 再导入。")
    raise RuntimeError(f"不支持的文件格式: {path.name}")


def _read_csv(path: Path) -> list[dict[str, Any]]:
    rows = _read_csv_raw(path)
    if not rows:
        return []
    headers = [clean_text(value) for value in rows[0]]
    data: list[dict[str, Any]] = []
    for row in rows[1:]:
        data.append(
            {
                normalize_key(headers[idx]): row[idx] if idx < len(row) else ""
                for idx in range(len(headers))
            }
        )
    return data


def _read_csv_raw(path: Path) -> list[list[Any]]:
    encodings = ("utf-8-sig", "utf-8", "gbk", "gb18030")
    last_error: Exception | None = None
    for encoding in encodings:
        try:
            with path.open("r", encoding=encoding, newline="") as file:
                return list(csv.reader(file))
        except Exception as exc:
            last_error = exc
    raise RuntimeError(f"CSV 读取失败: {path.name}: {last_error}") from last_error


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    raw_rows = _read_xlsx_raw(path)
    if not raw_rows:
        return []
    header_row = raw_rows[0]
    rows = raw_rows[1:]
    headers = [clean_text(value) for value in header_row]
    data: list[dict[str, Any]] = []
    for row in rows:
        data.append(
            {
                normalize_key(headers[idx]): (row[idx] if idx < len(row) else "")
                for idx in range(len(headers))
            }
        )
    return data


def _read_xlsx_raw(path: Path) -> list[list[Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    return [list(row) for row in rows]


def ensure_required_fields(rows: list[dict[str, Any]], fields: list[str], file_label: str) -> None:
    if not rows:
        raise RuntimeError(f"{file_label} 没有可读取的数据。")
    sample = rows[0]
    missing = [field for field in fields if field not in sample]
    if missing:
        raise RuntimeError(f"{file_label} 缺少字段: {', '.join(missing)}")


def build_sku_mapping(rows: list[dict[str, Any]]) -> dict[str, str]:
    ensure_required_fields(rows, SKU_MAPPING_REQUIRED_FIELDS, "运营对应SKU")
    mapping: dict[str, str] = {}
    for row in rows:
        sku = normalize_token(row.get("pcr_product_sku"))
        operator = clean_text(row.get("运营"))
        if sku:
            mapping[sku] = operator
    return mapping


def split_refs(value: Any) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    chunks = []
    current = []
    for char in text:
        if char in ",，;；\n\r\t ":
            if current:
                chunks.append("".join(current))
                current = []
        else:
            current.append(char)
    if current:
        chunks.append("".join(current))
    return [normalize_token(chunk) for chunk in chunks if normalize_token(chunk)]


def find_column_by_alias(columns: list[str], aliases: list[str]) -> str | None:
    alias_tokens = {normalize_token(alias) for alias in aliases}
    for column in columns:
        if normalize_token(column) in alias_tokens:
            return column
    return None


def find_header_row_index(raw_rows: list[list[Any]], required_alias_groups: list[list[str]]) -> int:
    normalized_groups = [{normalize_token(alias) for alias in group} for group in required_alias_groups]
    for index, row in enumerate(raw_rows):
        values = {normalize_token(value) for value in row if clean_text(value)}
        if values and all(group & values for group in normalized_groups):
            return index
    return -1


def rows_to_dicts(raw_rows: list[list[Any]], header_index: int) -> list[dict[str, Any]]:
    header_row = raw_rows[header_index]
    headers = [clean_text(value) for value in header_row]
    data: list[dict[str, Any]] = []
    for row in raw_rows[header_index + 1:]:
        if not any(clean_text(value) for value in row):
            continue
        data.append(
            {
                normalize_key(headers[idx]): (row[idx] if idx < len(row) else "")
                for idx in range(len(headers))
            }
        )
    return data


def _extract_order_rows_from_raw(
    raw_rows: list[list[Any]],
    file_label: str,
    allow_empty: bool = False,
) -> list[dict[str, Any]]:
    if not raw_rows:
        return []
    header_index = find_header_row_index(raw_rows, [[field] for field in ORDER_REQUIRED_FIELDS])
    if header_index == -1:
        if allow_empty:
            return []
        raise RuntimeError(f"{file_label} 未识别到订单表头。")
    rows = rows_to_dicts(raw_rows, header_index)
    if not rows and not allow_empty:
        raise RuntimeError(f"{file_label} 表头已识别，但没有可用订单数据。")
    return rows


def get_display_platform_value(platform_name: str, row: dict[str, Any]) -> str:
    normalized_platform = clean_text(platform_name).upper()
    if normalized_platform == "CD":
        return clean_text(row.get("user_account"))
    return clean_text(row.get("platform_user_name"))


def normalize_platform_name(value: Any) -> str:
    return clean_text(value).upper()


def build_operator_layout(operator_names: list[str]) -> dict[str, int]:
    left_col = 1
    middle_col = 11
    right_col = 21

    layout_columns: dict[str, int] = {
        "陆美婷": left_col,
        "陈嘉仪": left_col,
        "杨可依": left_col,
        "卢善程": middle_col,
        "郑逸君": middle_col,
        "许悦娜": right_col,
    }
    next_column_cycle = [middle_col, left_col]
    cycle_index = 0
    layout: dict[str, int] = {}

    for operator in operator_names:
        column = layout_columns.get(operator)
        if column is None:
            column = next_column_cycle[cycle_index % len(next_column_cycle)]
            cycle_index += 1
        layout[operator] = column
    return layout


def build_summary_table(grouped_rows: dict[str, list[dict[str, Any]]], summary_operators: list[str]) -> list[list[Any]]:
    summary_map: dict[str, dict[str, float]] = {}
    for operator in summary_operators:
        for row in grouped_rows.get(operator, []):
            platform = clean_text(row.get("平台"))
            if not platform or platform == "合计":
                continue
            if platform not in summary_map:
                summary_map[platform] = {
                    "销量": 0.0,
                    "销售额": 0.0,
                    "含税毛利": 0.0,
                    "未税毛利": 0.0,
                }
            summary_map[platform]["销量"] += parse_number(row.get("销量"))
            summary_map[platform]["销售额"] += parse_number(row.get("销售额"))
            summary_map[platform]["含税毛利"] += parse_number(row.get("含税毛利"))
            summary_map[platform]["未税毛利"] += parse_number(row.get("未税毛利"))

    rows: list[list[Any]] = [["", "当日总销量", "当日总销售额", "未税毛利率", "含税毛利率"]]
    total_qty = 0.0
    total_sales = 0.0
    total_tax_profit = 0.0
    total_no_tax_profit = 0.0

    for platform, data in summary_map.items():
        qty = round(data["销量"], 2)
        sales = round(data["销售额"], 2)
        total_qty += qty
        total_sales += sales
        total_tax_profit += data["含税毛利"]
        total_no_tax_profit += data["未税毛利"]
        rows.append([platform, qty, sales, "", ""])

    rows.append(
        [
            "合计",
            round(total_qty, 2),
            round(total_sales, 2),
            f"{(total_no_tax_profit / total_sales * 100):.2f}%" if total_sales else "",
            f"{(total_tax_profit / total_sales * 100):.2f}%" if total_sales else "",
        ]
    )
    return rows


def build_overall_totals(grouped_rows: dict[str, list[dict[str, Any]]]) -> tuple[float, float]:
    total_qty = 0.0
    total_sales = 0.0
    for rows in grouped_rows.values():
        for row in rows:
            total_qty += parse_number(row.get("销量"))
            total_sales += parse_number(row.get("销售额"))
    return round(total_sales, 2), round(total_qty, 2)


def build_cost_index(cost_report_path: Path, eur_to_cny_rate: float) -> list[dict[str, Any]]:
    raw_rows = read_raw_table(cost_report_path)
    if not raw_rows:
        return []
    header_index = find_header_row_index(
        raw_rows,
        [COST_ORDER_ALIASES, COST_SKU_ALIASES, COST_TOTAL_ALIASES],
    )
    if header_index == -1:
        raise RuntimeError("成本报表未识别到订单号、SKU 或总成本表头。")
    cost_rows = rows_to_dicts(raw_rows, header_index)
    if not cost_rows:
        raise RuntimeError("成本报表表头已识别，但没有可用数据行。")
    columns = list(cost_rows[0].keys())
    order_key = find_column_by_alias(columns, COST_ORDER_ALIASES)
    sku_key = find_column_by_alias(columns, COST_SKU_ALIASES)
    total_key = find_column_by_alias(columns, COST_TOTAL_ALIASES)
    unit_key = find_column_by_alias(columns, COST_UNIT_ALIASES)
    cny_total_key = None
    if not total_key:
        cny_total_key = find_column_by_alias(columns, ["总成本（人民币）", "总成本", "__EMPTY_11"])
    qty_key = find_column_by_alias(columns, ["总数量", "数量", "qty", "__EMPTY_10"])

    if not order_key or not sku_key:
        raise RuntimeError("成本报表未识别到订单号或 SKU 字段。")
    if not total_key and not cny_total_key:
        raise RuntimeError("成本报表未识别到总成本字段。")

    normalized_rows: list[dict[str, Any]] = []
    for row in cost_rows:
        order_no = normalize_token(row.get(order_key))
        sku_detail = normalize_token(row.get(sku_key))
        total_cost = parse_number(row.get(total_key)) if total_key else 0.0
        if not total_cost and cny_total_key:
            total_cost = round(parse_number(row.get(cny_total_key)) / eur_to_cny_rate, 2)
        unit_cost = parse_number(row.get(unit_key)) if unit_key else 0.0
        if not unit_cost and qty_key:
            qty = parse_number(row.get(qty_key))
            if qty:
                unit_cost = round(total_cost / qty, 2)
        normalized_rows.append(
            {
                "order_no": order_no,
                "sku_detail": sku_detail,
                "total_cost": total_cost,
                "unit_cost": unit_cost,
            }
        )
    return normalized_rows


def match_cost(cost_rows: list[dict[str, Any]], reference_no: Any, sku: Any, qty: float) -> float | None:
    refs = split_refs(reference_no)
    sku_token = normalize_token(sku)
    candidates: list[dict[str, Any]] = []
    for ref in refs:
        candidates.extend([row for row in cost_rows if row["order_no"] == ref])

    matched = None
    if candidates:
        matched = next((row for row in candidates if sku_token and sku_token in row["sku_detail"]), candidates[0])
    elif sku_token:
        matched = next(
            (
                row
                for row in cost_rows
                if row["sku_detail"] == sku_token or sku_token in row["sku_detail"]
            ),
            None,
        )

    if not matched:
        return None
    unit_cost = matched["unit_cost"]
    total_cost = matched["total_cost"]
    if unit_cost:
        return round(unit_cost * (qty or 1), 2)
    if total_cost:
        return round(total_cost, 2)
    return None


def process_daily_sales(
    platform_configs: list[dict[str, Any]],
    platform_files: dict[str, Path | str],
    sku_mapping_path: Path,
    cost_report_path: Path,
    output_path: Path,
    eur_to_cny_rate: float,
) -> ProcessingResult:
    sku_mapping_rows = read_table(sku_mapping_path)
    sku_mapping = build_sku_mapping(sku_mapping_rows)

    cost_rows = build_cost_index(cost_report_path, eur_to_cny_rate)
    grouped: dict[str, dict[tuple[str, str], dict[str, Any]]] = defaultdict(dict)
    unmatched_rows: list[dict[str, str]] = []

    for platform_config in platform_configs:
        platform_name = clean_text(platform_config["name"])
        normalized_platform = normalize_platform_name(platform_name)
        operator_name = clean_text(platform_config["operator"])
        requires_cost = bool(platform_config["requires_cost"])
        platform_file = platform_files[platform_name]
        if platform_file == "__NO_ORDER__":
            continue
        rows = read_order_tables(platform_file)
        ensure_required_fields(rows, ORDER_REQUIRED_FIELDS, f"{platform_name} 文件")

        for row in rows:
            order_status = clean_text(row.get("order_status"))
            if order_status in FILTER_STATUSES:
                continue
            unit_price = parse_number(row.get("unit_price"))
            qty = parse_number(row.get("qty"))
            if not unit_price or not qty:
                continue

            sku = normalize_token(row.get("pcr_product_sku"))
            matched_operator = clean_text(sku_mapping.get(sku, "")) if normalized_platform in COST_MATCH_PLATFORMS else ""
            if normalized_platform in COST_MATCH_PLATFORMS and not matched_operator:
                unmatched_rows.append(
                    {
                        "platform": platform_name,
                        "user_account": clean_text(row.get("user_account") or row.get("platform_user_name")),
                        "refrence_no": clean_text(
                            row.get("refrence_no") or row.get("refrence_no_platform")
                        ),
                    }
                )
                continue

            default_operators = split_operator_names(operator_name)
            final_operator = matched_operator or (default_operators[0] if default_operators else operator_name)
            total_sales = round(unit_price * qty, 2)
            tax_profit: float | None = None
            no_tax_profit: float | None = None
            profit_eligible = not requires_cost

            if requires_cost:
                matched_cost = match_cost(cost_rows, row.get("refrence_no_platform"), row.get("pcr_product_sku"), qty)
                if matched_cost is not None:
                    cost_value = round(matched_cost, 2)
                    tax_profit = round(total_sales - cost_value - total_sales * 0.2, 2)
                    no_tax_profit = round(tax_profit + total_sales * 0.2, 2)
                    profit_eligible = True

            display_platform = get_display_platform_value(platform_name, row) or platform_name
            product_name = clean_text(row.get("warehouse_product_title"))
            group_key = (display_platform, product_name)
            existing = grouped[final_operator].get(group_key)
            if not existing:
                existing = {
                    "平台": display_platform,
                    "产品名称": product_name,
                    "销量": 0.0,
                    "销售额": 0.0,
                    "含税毛利": 0.0,
                    "未税毛利": 0.0,
                    "需要成本": requires_cost,
                    "可计算毛利": False,
                }
                grouped[final_operator][group_key] = existing

            existing["销量"] += qty
            existing["销售额"] += total_sales
            existing["可计算毛利"] = existing["可计算毛利"] or profit_eligible
            if requires_cost and tax_profit is not None and no_tax_profit is not None:
                existing["含税毛利"] += tax_profit
                existing["未税毛利"] += no_tax_profit

    grouped_rows: dict[str, list[dict[str, Any]]] = {}
    for operator, rows in grouped.items():
        items = list(rows.values())
        items.sort(key=lambda item: item["销量"], reverse=True)
        for item in items:
            sales = item["销售额"]
            qty = item["销量"]
            if item["需要成本"] and item["可计算毛利"]:
                item["含税毛利"] = round(item["含税毛利"], 2)
                item["未税毛利"] = round(item["未税毛利"], 2)
                item["含税毛利率"] = f"{(item['含税毛利'] / sales * 100):.2f}%" if sales else ""
                item["未税毛利率"] = f"{(item['未税毛利'] / sales * 100):.2f}%" if sales else ""
            else:
                item["含税毛利"] = ""
                item["未税毛利"] = ""
                item["含税毛利率"] = ""
                item["未税毛利率"] = ""
            item["销量"] = round(qty, 2)
            item["销售额"] = round(sales, 2)
            item["平均单价"] = round(sales / qty, 2) if qty else 0
        grouped_rows[operator] = items

    report_date = date.today() - timedelta(days=1)
    write_report(output_path, report_date, grouped_rows, platform_configs)
    return ProcessingResult(
        output_path=output_path,
        report_date=report_date,
        grouped_rows=grouped_rows,
        unmatched_rows=unmatched_rows,
    )


def write_report(
    output_path: Path,
    report_date: date,
    grouped_rows: dict[str, list[dict[str, Any]]],
    platform_configs: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "销售数据"

    sheet["A1"] = f"{report_date:%Y年%m月%d日} 销售数据"
    sheet["A1"].font = Font(bold=True, size=14)

    seen_operators: list[str] = []
    for config in platform_configs:
        for operator in split_operator_names(config["operator"]):
            if operator not in seen_operators:
                seen_operators.append(operator)

    headers = ["平台", "产品名称", "销量", "销售额", "含税毛利", "未税毛利", "含税毛利率", "未税毛利率", "平均单价"]
    layout = build_operator_layout(seen_operators)
    row_positions = {
        1: 3,
        11: 3,
        21: 3,
    }

    for operator in seen_operators:
        start_col = layout[operator]
        start_row = row_positions[start_col]
        rows = grouped_rows.get(operator, [])
        current_row = start_row

        table_headers = [operator, *headers[1:]]
        for header_offset, header in enumerate(table_headers):
            sheet.cell(row=current_row, column=start_col + header_offset, value=header).font = Font(bold=True)
        current_row += 1

        total_qty = 0.0
        total_sales = 0.0
        total_tax_profit = 0.0
        total_no_tax_profit = 0.0
        requires_cost = False

        for item in rows:
            values = [
                item["平台"],
                item["产品名称"],
                item["销量"],
                item["销售额"],
                item["含税毛利"],
                item["未税毛利"],
                item["含税毛利率"],
                item["未税毛利率"],
                item["平均单价"],
            ]
            for value_offset, value in enumerate(values):
                sheet.cell(row=current_row, column=start_col + value_offset, value=value)
            total_qty += parse_number(item["销量"])
            total_sales += parse_number(item["销售额"])
            if bool(item["需要成本"]) and bool(item.get("可计算毛利")):
                total_tax_profit += parse_number(item["含税毛利"])
                total_no_tax_profit += parse_number(item["未税毛利"])
                requires_cost = True
            current_row += 1

        total_values = [
            "合计",
            "",
            round(total_qty, 2),
            round(total_sales, 2),
            round(total_tax_profit, 2) if requires_cost else "",
            round(total_no_tax_profit, 2) if requires_cost else "",
            f"{(total_tax_profit / total_sales * 100):.2f}%" if requires_cost and total_sales else "",
            f"{(total_no_tax_profit / total_sales * 100):.2f}%" if requires_cost and total_sales else "",
            round(total_sales / total_qty, 2) if total_qty else "",
        ]
        for value_offset, value in enumerate(total_values):
            sheet.cell(row=current_row, column=start_col + value_offset, value=value)
        row_positions[start_col] = current_row + 2

    summary_operators: list[str] = []
    for config in platform_configs:
        if bool(config.get("requires_cost")):
            for operator in split_operator_names(config.get("operator")):
                if operator not in summary_operators:
                    summary_operators.append(operator)

    summary_rows = build_summary_table(grouped_rows, summary_operators)
    left_start_row = row_positions[1]
    for row_offset, row_values in enumerate(summary_rows):
        for col_offset, value in enumerate(row_values):
            sheet.cell(row=left_start_row + row_offset, column=1 + col_offset, value=value)

    total_sales, total_qty = build_overall_totals(grouped_rows)
    totals_row = left_start_row + len(summary_rows) + 1
    sheet.cell(row=totals_row, column=1, value="总金额：")
    sheet.cell(row=totals_row, column=2, value=total_sales)
    sheet.cell(row=totals_row + 1, column=1, value="总销量：")
    sheet.cell(row=totals_row + 1, column=2, value=total_qty)

    for column_cells in sheet.columns:
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(width + 2, 10), 30)

    workbook.save(output_path)
